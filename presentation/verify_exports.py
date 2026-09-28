"""Проверка только артефактов защиты: PPTX, Office PDF и эталонной геометрии.

Не запускает приложение, selector, MAX, Docker или сбор источников.
"""
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone
import argparse, hashlib, json, re, zipfile
import xml.etree.ElementTree as ET
import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from render import ROOT, REPO, W, H, load_source, digest, write_json
from export_pptx import speaker_notes


def normalize(s):
    return ' '.join(s.replace('\u2010', '-').split())


def compact(s):
    return re.sub(r'\s+', '', normalize(s))


def glyphs(page):
    # Office adds transparent accessibility labels to hyperlink annotations.
    # Check those separately; compare the visible glyphs and their positions.
    return [(normalize(chr(c[0])), c[2]) for span in page.get_texttrace()
            if span['opacity']>0 and span['type']!=3
            for c in span['chars'] if not chr(c[0]).isspace()]


def shapes(seq):
    for shape in seq:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes(shape.shapes)


def verify(pptx_dir, pdf_dir, vector_dir, report_dir):
    source=load_source(); facts=json.loads((REPO/'docs/DEFENSE_FACTS.json').read_text(encoding='utf8'))
    report_dir.mkdir(parents=True,exist_ok=True)
    office=json.loads((pdf_dir/'office.json').read_text(encoding='utf8'))
    assert all(x['textAndGroupSavedReopened'] for x in office['decks'])
    report={'renderer':office['renderer'], 'applicationCommit':facts['applicationCommit'],
            'datasetVersion':facts['dataset']['version'], 'measurementDate':facts['measurementDate'],
            'applicationTests':facts.get('currentChecks',facts['recordedChecks']),
            'scope':'Экспорт и редактируемость; проверки приложения записаны отдельно',
            'decks':{},'officeEditability':office['decks']}
    decks=[]; pdfs=[]; notes=speaker_notes(source)
    for name,count,offset in [('cultural-plan',13,0),('cultural-plan-submission-preview',14,1)]:
        pptx=pptx_dir/(name+'.pptx'); pdf=pdf_dir/(name+'.pdf')
        deck=Presentation(pptx); doc=fitz.open(pdf); vector=fitz.open(vector_dir/(name+'.pdf'))
        decks.append(deck); pdfs.append(doc)
        assert len(deck.slides)==len(doc)==len(vector)==count
        assert (deck.slide_width.pt,deck.slide_height.pt)==(W,H)
        assert not doc.embfile_count() and not doc.is_encrypted
        ids=(['T00'] if offset else [])+[s['id'] for s in source['slides']]
        counter=Counter(); link_count=0; max_shift=0.; text_frames=0
        with zipfile.ZipFile(pptx) as z:
            assert z.testzip() is None
            assert not any('/embeddings/' in p or p.endswith(('.ttf','.otf','.fntdata','.svg')) for p in z.namelist())
            assert not any(p.startswith('ppt/media/') for p in z.namelist()), 'В исходных PPTX не было изображений'
            xml='\n'.join(z.read(n).decode('utf8') for n in z.namelist() if n.endswith(('.xml','.rels')))
            assert not re.search(r'(?:Bearer\s+[A-Za-z0-9_.-]{20,}|BEGIN [A-Z ]*PRIVATE KEY|C:\\Users\\|BastaPC|file:///)',xml)
            for name_in_zip in z.namelist():
                if name_in_zip.endswith('.rels'):
                    for rel in ET.fromstring(z.read(name_in_zip)):
                        if rel.get('TargetMode')=='External':assert rel.get('Target').startswith('https://')
        for index,(slide,page,reference) in enumerate(zip(deck.slides,doc,vector)):
            assert tuple(page.rect)==(0,0,W,H)
            assert normalize(slide.notes_slide.notes_text_frame.text)==normalize(notes[ids[index]])
            visible=[]
            for shape in shapes(slide.shapes):
                counter[str(shape.shape_type)]+=1
                if shape.has_text_frame and shape.text:
                    text_frames+=1; visible.append(shape.text)
                    assert shape.left>=0 and shape.top>=0
                    assert shape.left+shape.width<=deck.slide_width+1270
                    assert shape.top+shape.height<=deck.slide_height+1270
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.text: assert run.font.name=='Segoe UI' and run.font.size.pt>=14
                address=shape.click_action.hyperlink.address if shape.shape_type!=MSO_SHAPE_TYPE.GROUP else None
                if address: assert address.startswith('https://')
            a,b=glyphs(page),glyphs(reference)
            actual=''.join(c for c,_ in a)
            assert compact('\n'.join(visible))==actual,(name,index,'PPTX/PDF text')
            assert compact(reference.get_text())==actual,(name,index,'vector/Office text')
            for span in page.get_texttrace():
                if span['opacity']==0 or span['type']==3:
                    label=''.join(chr(c[0]) for c in span['chars'])
                    assert any(compact(label) in compact(v) for v in visible),('unexpected invisible text',label)
            assert not any(t in actual for t in ['{{','\ufffd','257','62 события','пяти источников'])
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines',[]):
                    for span in line['spans']:
                        assert page.rect.contains(fitz.Rect(span['bbox'])),(name,index,span['text'])
            assert all('Segoe' in s['font'] for s in page.get_texttrace() if s['opacity']>0 and s['type']!=3)
            urls=sorted(x['uri'] for x in page.get_links())
            assert urls==sorted(x['uri'] for x in reference.get_links()),(name,index,'links')
            link_count+=len(urls)
            # Compare glyph positions across the established renderer and Office.
            # Text already agrees; span segmentation is allowed to differ.
            assert ''.join(c for c,_ in a)==''.join(c for c,_ in b)
            shift=max((max(abs(x[0]-y[0]),abs(x[1]-y[1])) for (_,x),(_,y) in zip(a,b)),default=0)
            max_shift=max(max_shift,shift)
            assert shift<2.0,(name,index,shift)
            page.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(report_dir/f'{name}-{index+1:02d}.png')
        assert counter[str(MSO_SHAPE_TYPE.TEXT_BOX)]>100 and counter[str(MSO_SHAPE_TYPE.GROUP)]>10
        assert counter[str(MSO_SHAPE_TYPE.LINE)]>10 and not counter[str(MSO_SHAPE_TYPE.PICTURE)]
        for start in range(0,count,4):
            sheet=fitz.open(); p=sheet.new_page(width=1456,height=826)
            for j,index in enumerate(range(start,min(start+4,count))):
                x=8+(j%2)*728; y=8+(j//2)*409
                p.show_pdf_page(fitz.Rect(x,y,x+720,y+405),doc,index)
            p.get_pixmap().save(report_dir/f'{name}-contact-{start//4+1}.png'); sheet.close()
        report['decks'][name]={'slides':count,'nativeShapes':dict(counter),'textFrames':text_frames,
                'links':link_count,'allNotesMatch':True,'pptxPdfTextMatch':True,'vectorPdfTextMatch':True,
                'maxGlyphOriginDifferencePt':round(max_shift,4),
                'files':{pptx.name:digest(pptx),pdf.name:digest(pdf)}}
        vector.close()
    for i in range(13):
        a,b=pdfs[0][i],pdfs[1][i+1]
        assert a.get_text()==b.get_text()
        assert a.get_pixmap().samples==b.get_pixmap().samples,('edition pixels',i)
        assert decks[0].slides[i].shapes._spTree.xml==decks[1].slides[i+1].shapes._spTree.xml,('native shapes',i)
    report['identicalProductPages']=13
    assert facts['applicationCommit'] in pdfs[1][0].get_text()
    assert all(n in pdfs[0][0].get_text() for n in source['meta']['participants'])
    assert all(v['name'] in pdfs[0][8].get_text() for v in facts['cities'].values())
    assert '₽' in ''.join(p.get_text() for d in pdfs for p in d)
    for d in pdfs:d.close()
    write_json(report_dir/'validation.json',report)
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    for arg in ['pptx-dir','pdf-dir','vector-dir','report-dir']:p.add_argument('--'+arg,type=Path,required=True)
    p.add_argument('--write-build-info',action='store_true',help='Записать метаданные проверенных финальных файлов')
    a=p.parse_args(); result=verify(a.pptx_dir,a.pdf_dir,a.vector_dir,a.report_dir)
    if a.write_build_info:
        outputs={k:v for d in result['decks'].values() for k,v in d['files'].items()}
        for name,expected in outputs.items():
            assert digest(ROOT/name)==expected,'Сначала скопируйте проверенные файлы в presentation/: '+name
        inputs=['slides.json','render.py','export_pptx.py','export_office.ps1','verify_exports.py',
                'SPEAKER_NOTES.md','requirements.txt','requirements-pptx.txt','../docs/DEFENSE_FACTS.json']
        build={'reviewed_application_commit':result['applicationCommit'],'dataset_version':result['datasetVersion'],
            'measurement_date':result['measurementDate'],'artifact_metadata_updated_at':datetime.now(timezone.utc).isoformat(),
            'renderer':result['renderer'],'python':'3.12','pymupdf':fitz.VersionBind,
            'inputs_sha256':{n:digest(ROOT/n) for n in inputs},'outputs_sha256':outputs,
            'pages':{'product':13,'submission':14},'identical_product_pages':13,
            'fonts':'Segoe UI Regular/Bold с хоста; font-файлы не включены',
            'note':'Время генерации не меняет дату измерения или свежесть фактов.'}
        write_json(ROOT/'BUILD_INFO.json',build)
